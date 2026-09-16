


import os
import json
import traceback

import pandas as pd
from dotenv import load_dotenv
from docx import Document

from src.mcqgenerator.utills import read_file, get_table_data

import streamlit as st
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.prompts import PromptTemplate
from src.mcqgenerator.logger import logging


load_dotenv()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


st.title("MCQ Generator using Gemini and Langchain")


# Store generated MCQs
if "mcq_text" not in st.session_state:
    st.session_state.mcq_text = None


# -----------------------------
# FORM
# -----------------------------

with st.form("my_form"):

    uploaded_file = st.file_uploader(
        "Upload a PDF or TXT file",
        type=["pdf", "txt"]
    )

    number = st.number_input(
        "Number of MCQs to generate",
        min_value=1,
        max_value=50,
        value=3,
        step=1,
        key="num_mcqs"
    )

    subject = st.text_input(
        "Enter the subject for the MCQs",
        value="General Knowledge",
        key="subject"
    )

    tone = st.selectbox(
        "Select the tone for the MCQs",
        ["Formal", "Informal", "Humorous", "Serious"],
        key="tone"
    )

    submit_button = st.form_submit_button(
        label="Generate MCQs"
    )


    # Generate MCQs
    if submit_button and uploaded_file is not None:

        with st.spinner("Generating MCQs..."):

            try:

                # Read file
                text = read_file(uploaded_file)

                num_mcqs = st.session_state.num_mcqs
                subject = st.session_state.subject
                tone = st.session_state.tone


                # Gemini
                llm = ChatGoogleGenerativeAI(
                    model="gemini-3.5-flash",
                    google_api_key=os.getenv("genai_api_key"),
                    temperature=0.5
                )


                # Prompt
                prompt_template = PromptTemplate(
                    input_variables=[
                        "text",
                        "num_mcqs",
                        "subject",
                        "tone"
                    ],
                    template="""
You are an expert in generating multiple choice questions (MCQs)
based on the provided text.

Generate {num_mcqs} MCQs on the subject of {subject}
with a {tone} tone.

For each question provide:

Question
A) Option A
B) Option B
C) Option C
D) Option D

Correct Answer
Explanation

Text:
{text}
"""
                )


                prompt = prompt_template.format(
                    text=text,
                    num_mcqs=num_mcqs,
                    subject=subject,
                    tone=tone
                )


                # Generate response
                response = llm.invoke(prompt)


                # Extract response text
                content = response.content

                if isinstance(content, list):

                    response_text = ""

                    for item in content:

                        if (
                            isinstance(item, dict)
                            and item.get("type") == "text"
                        ):
                            response_text += item.get("text", "")

                else:

                    response_text = str(content)


                # Save MCQs in session state
                st.session_state.mcq_text = response_text


            except Exception as e:

                logging.error(
                    f"Error generating MCQs: {e}"
                )

                st.error(
                    f"Error generating MCQs: {e}"
                )


# -----------------------------
# OUTSIDE FORM
# -----------------------------

if st.session_state.mcq_text:

    # Display MCQs
    st.markdown("### Generated MCQs")

    st.markdown(
        st.session_state.mcq_text
    )


    # -----------------------------
    # CREATE WORD FILE
    # -----------------------------

    doc = Document()

    doc.add_heading(
        "Generated MCQs",
        level=1
    )

    doc.add_paragraph(
        st.session_state.mcq_text
    )


    word_file = os.path.join(
        BASE_DIR,
        "generated_mcqs.docx"
    )

    doc.save(word_file)


    # -----------------------------
    # DOWNLOAD BUTTON
    # -----------------------------

    with open(word_file, "rb") as file:

        st.download_button(
            label="Download MCQs as Word",
            data=file,
            file_name="generated_mcqs.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )